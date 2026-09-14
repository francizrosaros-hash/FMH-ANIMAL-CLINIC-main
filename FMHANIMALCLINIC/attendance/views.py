"""Views for Attendance and Biometrics Management."""
from django.shortcuts import render, redirect, get_object_or_404
from django.http import FileResponse, Http404
from django.http import HttpResponse
from django.core.files.storage import default_storage
from django.views.decorators.http import require_http_methods
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db.models import Q, Count, Sum
from django.utils import timezone
from calendar import monthrange
from datetime import date, datetime, timedelta
from decimal import Decimal
import logging
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accounts.decorators import admin_only
from employees.models import StaffMember, VetSchedule
from branches.models import Branch
from .models import DailyAttendance, MonthlyAttendanceSummary, AttendanceUpload
from .forms import (
    AttendanceImportForm,
    DailyAttendanceForm,
    AttendanceFilterForm,
)
from .services import AttendanceImportService, AttendanceProcessor
from settings.utils import get_setting
from payroll.models import PayrollPeriod, Payslip


def _system_daily_salary(staff):
    """
    Calculate daily salary from the staff base salary and actual working days.
    Uses working days from the current month's attendance data if available,
    otherwise falls back to calculating working days (Mon-Fri) for the current month.
    """
    if not staff.salary:
        return Decimal('0')
    
    today = date.today()
    
    # Try to get actual working days from attendance summary for current month
    current_month_summary = MonthlyAttendanceSummary.objects.filter(
        staff=staff,
        period_start__year=today.year,
        period_start__month=today.month,
    ).first()
    
    if current_month_summary and current_month_summary.working_days > 0:
        return staff.salary / Decimal(str(current_month_summary.working_days))
    
    # Fallback: Calculate working days (Mon-Fri) for the current month
    from calendar import monthrange
    _, days_in_month = monthrange(today.year, today.month)
    working_days = 0
    for day in range(1, days_in_month + 1):
        check_date = date(today.year, today.month, day)
        if check_date.weekday() < 5:  # Monday=0, Friday=4
            working_days += 1
    
    if working_days > 0:
        return staff.salary / Decimal(str(working_days))
    
    return Decimal('0')

logger = logging.getLogger(__name__)


def can_import_attendance_for_month(year, month):
    """Allow re-imports until both payroll halves for the month are generated/released."""
    completed_statuses = [PayrollPeriod.Status.GENERATED, PayrollPeriod.Status.RELEASED]
    completed_half_types = set(
        PayrollPeriod.objects.filter(
            year=year,
            month=month,
            status__in=completed_statuses,
            period_type__in=[
                PayrollPeriod.PeriodType.SEMI_FIRST,
                PayrollPeriod.PeriodType.SEMI_SECOND,
            ],
        ).values_list('period_type', flat=True)
    )
    return not {
        PayrollPeriod.PeriodType.SEMI_FIRST,
        PayrollPeriod.PeriodType.SEMI_SECOND,
    }.issubset(completed_half_types)


def _delete_attendance_upload_files(year, month, upload=None):
    """Delete stored attendance files in the period and actual upload folders."""
    directories = {f'attendance/uploads/{year:04d}/{month:02d}'}
    if upload and upload.source_file:
        source_name = upload.source_file.name.replace('\\', '/')
        if '/' in source_name:
            directories.add(source_name.rsplit('/', 1)[0])

    for directory in directories:
        try:
            subdirectories, files = default_storage.listdir(directory)
            for filename in files:
                default_storage.delete(f'{directory}/{filename}')
            for subdirectory in subdirectories:
                nested_directory = f'{directory}/{subdirectory}'
                _, nested_files = default_storage.listdir(nested_directory)
                for filename in nested_files:
                    default_storage.delete(f'{nested_directory}/{filename}')
        except (FileNotFoundError, OSError):
            continue


def _attendance_role_label(staff):
    """Use the assigned RBAC role for attendance displays when available."""
    assigned_role = getattr(getattr(staff, 'user', None), 'assigned_role', None)
    return assigned_role.name if assigned_role else staff.get_position_display()


# ─────────────────── DASHBOARD ───────────────────
@login_required
@admin_only
def attendance_dashboard(request):
    """Attendance import dashboard for payroll."""
    return render(request, 'attendance/dashboard.html')


# ─────────────────── ATTENDANCE IMPORT ───────────────────
@login_required
@admin_only
def attendance_import(request):
    """Import one monthly summary attendance file for payroll."""
    if request.method == 'POST':
        form = AttendanceImportForm(request.POST, request.FILES)
        if form.is_valid():
            try:
                service = AttendanceImportService()
                records = form.get_records()
                unmatched_ids = AttendanceImportService.get_unmatched_biometric_ids(records)

                report_start = next(
                    (service.parse_date_value(record.get('Report Start')) for record in records
                     if isinstance(record, dict) and record.get('Report Start')),
                    None,
                )
                report_end = next(
                    (service.parse_date_value(record.get('Report End')) for record in records
                     if isinstance(record, dict) and record.get('Report End')),
                    None,
                )
                if not report_start:
                    attendance_dates = [
                        service.parse_date_value(service._get_first_matching_value(
                            record, {'date', 'attendance date', 'work date'}
                        ))
                        for record in records if isinstance(record, dict)
                    ]
                    attendance_dates = [attendance_date for attendance_date in attendance_dates if attendance_date]
                    if attendance_dates:
                        report_start = min(attendance_dates).replace(day=1)
                        report_end = report_start.replace(
                            day=monthrange(report_start.year, report_start.month)[1]
                        )

                if report_start and not can_import_attendance_for_month(report_start.year, report_start.month):
                    messages.error(
                        request,
                        f'Attendance for {report_start:%B %Y} cannot be imported because both the first-half and second-half payrolls for this month have already been generated or released.'
                    )
                    return redirect('attendance:import')

                existing_upload = None
                if report_start and report_end:
                    existing_upload = AttendanceUpload.objects.filter(
                        period_start__year=report_start.year,
                        period_start__month=report_start.month,
                    ).order_by('-period_end').first()
                    if existing_upload:
                        old_daily_deleted, old_summary_deleted = AttendanceImportService.clear_monthly_import_data(
                            report_start, report_end
                        )
                        logger.info(
                            f'Deleted {old_daily_deleted} old DailyAttendance and '
                            f'{old_summary_deleted} old MonthlyAttendanceSummary records '
                            f'for {report_start.strftime("%B %Y")}'
                        )
                    _delete_attendance_upload_files(report_start.year, report_start.month, existing_upload)

                imported, matched, errors = service.import_summary_records(
                    records,
                    source_filename=form.cleaned_data['import_file'].name,
                    uploaded_by=request.user,
                    skip_duplicates=not bool(existing_upload),
                )

                if unmatched_ids:
                    messages.warning(
                        request,
                        'The following biometric IDs were not matched to active staff: ' + ', '.join(unmatched_ids)
                    )

                if report_start and report_end:
                    import_file = form.cleaned_data['import_file']
                    import_file.seek(0)

                    if existing_upload:
                        existing_upload.period_start = report_start
                        existing_upload.period_end = report_end
                        existing_upload.source_file = import_file
                        existing_upload.source_filename = import_file.name
                        existing_upload.unmatched_biometric_ids = unmatched_ids
                        existing_upload.uploaded_by = request.user
                        existing_upload.save()
                    else:
                        AttendanceUpload.objects.create(
                            period_start=report_start,
                            period_end=report_end,
                            source_file=import_file,
                            source_filename=import_file.name,
                            unmatched_biometric_ids=unmatched_ids,
                            uploaded_by=request.user,
                        )

                    editable_periods = PayrollPeriod.objects.filter(
                        year=report_start.year,
                        month=report_start.month,
                        status__in=[PayrollPeriod.Status.DRAFT, PayrollPeriod.Status.EDITED],
                    )
                    for payroll_period in editable_periods:
                        for payslip in Payslip.objects.filter(payroll_period=payroll_period):
                            payslip.generate_from_employee()
                            payslip.save()
                
                message = (
                    f'Monthly import completed: {imported} records imported, '
                    f'{matched} matched to staff, {errors} errors.'
                )
                messages.success(request, message)
                logger.info(message)
                
                if records and all(record.get('Report Start') for record in records if isinstance(record, dict)):
                    report_month = service.parse_date_value(records[0].get('Report Start'))
                    if report_month:
                        return redirect(
                            f'/attendance/review/?year={report_month.year}&month={report_month.month}'
                        )
                return redirect('attendance:import')
                
            except Exception as e:
                logger.error(f'Import error: {str(e)}')
                messages.error(request, f'Import failed: {str(e)}')
    else:
        form = AttendanceImportForm()

    context = {
        'form': form,
        'page_title': 'Import Attendance Data',
    }
    
    return render(request, 'attendance/import.html', context)


@login_required
@admin_only
def attendance_upload_history(request):
    """List previously imported attendance files and their review state."""
    upload_history = AttendanceUpload.objects.select_related('uploaded_by').all()
    return render(request, 'attendance/history.html', {
        'upload_history': upload_history,
        'page_title': 'Attendance Upload History',
    })


@login_required
@admin_only
@require_http_methods(['POST'])
def attendance_import_delete(request, year, month):
    """
    Delete an imported month so a corrected file can be uploaded.
    
    Deletes:
    - All DailyAttendance records for the month
    - MonthlyAttendanceSummary records
    - AttendanceUpload record and source file
    
    Safety:
    - Prevents deletion if attendance is LOCKED (approved and finalized)
    - Prevents deletion if payroll is RELEASED (for audit trail)
    - Payslips remain unaffected (stored separately)
    """
    from payroll.models import PayrollPeriod, Payslip
    
    upload = AttendanceUpload.objects.filter(
        period_start__year=year,
        period_start__month=month,
    ).first()
    
    summaries = MonthlyAttendanceSummary.objects.filter(
        period_start__year=year,
        period_start__month=month,
    )
    
    daily_records = DailyAttendance.objects.filter(
        attendance_date__year=year,
        attendance_date__month=month,
    )
    
    # Check if attendance is locked
    if summaries.filter(review_status=MonthlyAttendanceSummary.ReviewStatus.LOCKED).exists():
        messages.error(request, 'Locked attendance cannot be deleted. It must be unlocked first.')
        return redirect('attendance:history')
    
    # Check if payroll was released for this month
    released_payroll = PayrollPeriod.objects.filter(
        year=year,
        month=month,
        status=PayrollPeriod.Status.RELEASED,
    ).exists()
    
    if released_payroll:
        messages.error(
            request,
            f'Cannot delete attendance for {month:02d}/{year} - payroll has been released. '
            'If you need to correct attendance, contact administration to reverse the payroll release first.'
        )
        return redirect('attendance:history')
    
    # If no data, inform user
    if not upload and not summaries.exists() and not daily_records.exists():
        messages.error(request, 'No attendance data exists for this month.')
        return redirect('attendance:history')
    
    # Delete all traces of the import
    daily_count, _ = daily_records.delete()
    summary_count, _ = summaries.delete()

    upload_storage_name = upload.source_file.name if upload and upload.source_file else None
    if upload:
        upload.delete()

    if upload_storage_name:
        upload.source_file.name = upload_storage_name
    _delete_attendance_upload_files(year, month, upload)
    
    logger.info(
        f'Deleted attendance for {month:02d}/{year}: '
        f'{daily_count} daily records, {summary_count} monthly summaries'
    )
    
    messages.success(
        request,
        f'Deleted all attendance data for {month:02d}/{year} '
        f'({daily_count} daily records, {summary_count} summary). '
        'Payroll records remain unaffected.'
    )
    return redirect('attendance:history')


@login_required
@admin_only
def attendance_import_file(request, upload_id):
    """Open the original uploaded attendance export."""
    upload = get_object_or_404(AttendanceUpload, id=upload_id)
    if not upload.source_file:
        raise Http404('The original attendance file is unavailable.')
    return FileResponse(upload.source_file.open('rb'), as_attachment=False, filename=upload.source_filename)


# ─────────────────── ATTENDANCE REVIEW ───────────────────
@login_required
@admin_only
def attendance_review(request):
    """Review and approve imported monthly attendance summaries."""
    year = _attendance_query_int(request.GET.get('year'), timezone.now().year, 2020, 2100)
    month = _attendance_query_int(request.GET.get('month'), timezone.now().month, 1, 12)
    summaries = list(MonthlyAttendanceSummary.objects.filter(
        period_start__year=year,
        period_start__month=month,
    ).select_related('staff', 'uploaded_by', 'approved_by').order_by('staff__last_name', 'staff__first_name'))
    for summary in summaries:
        summary.system_daily_salary = _system_daily_salary(summary.staff)

    total_records = len(summaries)
    context = {
        'monthly_summaries': summaries,
        'selected_year': year,
        'selected_month': month,
        'total_records': total_records,
        'page_title': 'Review Monthly Attendance',
        'attendance_upload': AttendanceUpload.objects.filter(
            period_start__year=year,
            period_start__month=month,
        ).first(),
    }
    return render(request, 'attendance/review.html', context)


@login_required
@admin_only
def attendance_edit(request, attendance_id):
    """Edit a daily attendance record."""
    attendance = get_object_or_404(DailyAttendance, id=attendance_id)
    
    if request.method == 'POST':
        form = DailyAttendanceForm(request.POST, instance=attendance)
        if form.is_valid():
            attendance = form.save(commit=False)
            attendance.is_manually_adjusted = True
            attendance.adjusted_by = request.user
            attendance.save()
            messages.success(request, 'Attendance record updated.')
            return redirect('attendance:review')
    else:
        form = DailyAttendanceForm(instance=attendance)
    
    # Get schedule for reference
    schedule = VetSchedule.objects.filter(
        staff=attendance.staff,
        date=attendance.attendance_date,
        is_available=True
    ).first()
    
    context = {
        'form': form,
        'attendance': attendance,
        'schedule': schedule,
        'page_title': f'Edit Attendance: {attendance.staff.full_name}',
    }
    
    return render(request, 'attendance/attendance_edit.html', context)


@login_required
@admin_only
@require_http_methods(["POST"])
def attendance_approve(request, attendance_id):
    """Approve an attendance record."""
    attendance = get_object_or_404(DailyAttendance, id=attendance_id)
    attendance.is_approved = True
    attendance.approved_by = request.user
    attendance.approved_at = timezone.now()
    attendance.save()
    messages.success(request, f'Attendance for {attendance.staff.full_name} approved.')
    return redirect('attendance:review')

# ─────────────────── REPORTS ───────────────────
def _attendance_query_int(value, default, minimum, maximum):
    """Parse formatted report query values without allowing invalid dates."""
    try:
        parsed = int(str(value).replace(',', '').strip())
    except (TypeError, ValueError):
        return default
    return parsed if minimum <= parsed <= maximum else default


@login_required
@admin_only
def attendance_summary(request):
    """Monthly attendance summary report."""
    # Get current month or specified month.
    current_year = timezone.now().year
    year = _attendance_query_int(request.GET.get('year'), current_year, 2020, 2100)
    month = _attendance_query_int(request.GET.get('month'), timezone.now().month, 1, 12)
    branch_id = request.GET.get('branch') or ''
    
    # Calculate date range
    start_date = timezone.datetime(year, month, 1).date()
    if month == 12:
        end_date = timezone.datetime(year + 1, 1, 1).date() - timedelta(days=1)
    else:
        end_date = timezone.datetime(year, month + 1, 1).date() - timedelta(days=1)
    
    # Get staff and optionally scope the report to a branch.
    staff_members = StaffMember.objects.filter(
        is_active=True,
    ).exclude(
        user__assigned_role__code='superadmin'
    ).select_related('user', 'user__assigned_role')
    if branch_id:
        staff_members = staff_members.filter(branch_id=branch_id)
    staff_members = staff_members.order_by('last_name', 'first_name')
    
    # Build summary
    summary_data = []
    total_scheduled_days = 0
    total_present_days = 0
    total_overtime_minutes = 0
    for staff in staff_members:
        monthly_summary = MonthlyAttendanceSummary.objects.filter(
            staff=staff,
            period_start__year=year,
            period_start__month=month,
        ).order_by('-period_end').first()
        attendance_records = DailyAttendance.objects.filter(
            staff=staff,
            attendance_date__range=[start_date, end_date],
        )
        
        present_days = monthly_summary.attendance_days if monthly_summary else attendance_records.filter(is_present=True).count()
        absent_days = monthly_summary.absence_days if monthly_summary else attendance_records.filter(is_present=False).count()
        late_days = monthly_summary.late_days if monthly_summary else attendance_records.filter(status='LATE').count()
        total_late_minutes = attendance_records.aggregate(
            total=Sum('late_minutes')
        )['total'] or 0
        staff_overtime_minutes = (
            monthly_summary.overtime_hours * Decimal('60')
            if monthly_summary else attendance_records.aggregate(
                total=Sum('overtime_minutes')
            )['total'] or 0
        )

        scheduled_days = monthly_summary.working_days if monthly_summary else VetSchedule.objects.filter(
            staff=staff,
            date__range=[start_date, end_date],
            is_available=True,
        ).count()
        denominator = scheduled_days or (present_days + absent_days)
        attendance_rate = (present_days / denominator * 100) if denominator else 0
        total_scheduled_days += scheduled_days
        total_present_days += present_days
        total_overtime_minutes += staff_overtime_minutes

        summary_data.append((staff, {
            'working_days': scheduled_days,
            'days_present': present_days,
            'days_absent': absent_days,
            'days_late': late_days,
            'total_late_minutes': total_late_minutes,
            'total_overtime_hours': Decimal(staff_overtime_minutes) / Decimal('60'),
            'sick_hours': monthly_summary.sick_hours if monthly_summary else Decimal('0'),
            'leave_hours': monthly_summary.leave_hours if monthly_summary else Decimal('0'),
            'daily_salary': _system_daily_salary(staff) if monthly_summary else Decimal('0'),
            'overtime_pay': monthly_summary.overtime_pay if monthly_summary else Decimal('0'),
            'charges': monthly_summary.charges if monthly_summary else Decimal('0'),
            'attendance_rate': attendance_rate,
        }))

    total_working_days = total_scheduled_days
    avg_attendance_rate = (
        total_present_days / total_scheduled_days * 100
        if total_scheduled_days else 0
    )
    
    context = {
            'summary_data': summary_data,
        'year': year,
        'month': month,
        'start_date': start_date,
        'end_date': end_date,
        'year_choices': range(current_year - 2, current_year + 2),
        'month_choices': [(index, timezone.datetime(2000, index, 1).strftime('%B')) for index in range(1, 13)],
        'branches': Branch.objects.filter(is_active=True),
        'selected_year': year,
        'selected_month': month,
        'selected_branch': int(branch_id) if branch_id.isdigit() else '',
        'total_working_days': total_working_days,
        'avg_attendance_rate': avg_attendance_rate,
        'total_overtime_hours': Decimal(total_overtime_minutes) / Decimal('60'),
        'page_title': f'Attendance Summary - {start_date.strftime("%B %Y")}',
    }
    
    return render(request, 'attendance/summary.html', context)


@login_required
@admin_only
def attendance_summary_excel(request):
    """Export the selected attendance summary as a formatted Excel workbook."""
    current_year = timezone.now().year
    year = _attendance_query_int(request.GET.get('year'), current_year, 2020, 2100)
    month = _attendance_query_int(request.GET.get('month'), timezone.now().month, 1, 12)
    branch_id = request.GET.get('branch') or ''
    start_date = date(year, month, 1)
    end_date = date(year + 1, 1, 1) - timedelta(days=1) if month == 12 else date(year, month + 1, 1) - timedelta(days=1)

    staff_members = StaffMember.objects.filter(is_active=True).exclude(
        user__assigned_role__code='superadmin'
    ).select_related('user', 'user__assigned_role')
    if branch_id.isdigit():
        staff_members = staff_members.filter(branch_id=branch_id)
    staff_members = staff_members.order_by('last_name', 'first_name')

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = 'Attendance Summary'
    worksheet.merge_cells('A1:N1')
    worksheet['A1'] = f'Attendance Summary - {start_date.strftime("%B %Y")}'
    worksheet['A1'].font = Font(bold=True, size=16, color='FFFFFF')
    worksheet['A1'].fill = PatternFill('solid', fgColor='00796B')
    worksheet['A1'].alignment = Alignment(horizontal='center')
    worksheet.merge_cells('A2:O2')
    worksheet['A2'] = f'Period: {start_date:%B %d, %Y} - {end_date:%B %d, %Y}'
    worksheet['A2'].font = Font(italic=True, color='5E6278')

    headers = [
        'Staff Member', 'Designation', 'Working Days', 'Days Present',
        'Days Absent', 'Days Late', 'Late Minutes', 'Overtime Hours',
        'Sick Hours', 'Leave Hours', 'Daily Salary', 'Overtime Pay',
        'Charges', 'Attendance Rate',
    ]
    header_row = 4
    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(header_row, column, header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='009688')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for row_index, staff in enumerate(staff_members, header_row + 1):
        monthly_summary = MonthlyAttendanceSummary.objects.filter(
            staff=staff, period_start__year=year, period_start__month=month,
        ).order_by('-period_end').first()
        attendance_records = DailyAttendance.objects.filter(
            staff=staff, attendance_date__range=[start_date, end_date],
        )
        present_days = monthly_summary.attendance_days if monthly_summary else attendance_records.filter(is_present=True).count()
        absent_days = monthly_summary.absence_days if monthly_summary else attendance_records.filter(is_present=False).count()
        late_days = monthly_summary.late_days if monthly_summary else attendance_records.filter(status='LATE').count()
        late_minutes = attendance_records.aggregate(total=Sum('late_minutes'))['total'] or 0
        overtime_hours = monthly_summary.overtime_hours if monthly_summary else Decimal(attendance_records.aggregate(total=Sum('overtime_minutes'))['total'] or 0) / Decimal('60')
        working_days = monthly_summary.working_days if monthly_summary else VetSchedule.objects.filter(staff=staff, date__range=[start_date, end_date], is_available=True).count()
        attendance_rate = present_days / working_days if working_days else 0
        values = [
            staff.full_name, _attendance_role_label(staff), working_days, present_days,
            absent_days, late_days, late_minutes, overtime_hours,
            monthly_summary.sick_hours if monthly_summary else 0,
            monthly_summary.leave_hours if monthly_summary else 0,
            _system_daily_salary(staff) if monthly_summary else 0,
            monthly_summary.overtime_pay if monthly_summary else 0,
            monthly_summary.charges if monthly_summary else 0,
            attendance_rate,
        ]
        for column, value in enumerate(values, 1):
            worksheet.cell(row_index, column, value)

    worksheet.freeze_panes = 'A5'
    worksheet.auto_filter.ref = f'A4:N{max(4, worksheet.max_row)}'
    worksheet.row_dimensions[1].height = 28
    worksheet.row_dimensions[4].height = 34
    widths = [24, 20, 14, 14, 14, 12, 14, 16, 12, 12, 15, 15, 14, 16]
    for column, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    for row in worksheet.iter_rows(min_row=5, max_row=worksheet.max_row, min_col=3, max_col=14):
        for cell in row:
            cell.alignment = Alignment(vertical='center')
            if cell.column in (11, 12, 13, 14):
                cell.number_format = '#,##0.00'
            elif cell.column in (8, 9, 10):
                cell.number_format = '0.00'
            elif cell.column == 16:
                cell.number_format = '0.0%'

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="attendance-summary-{year}-{month:02d}.xlsx"'
    workbook.save(response)
    return response
